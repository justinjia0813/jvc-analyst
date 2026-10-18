#!/usr/bin/env python3
"""Build three populated, synthetic examples and independent sensitivity expectations."""
import importlib.util
import json
from pathlib import Path

from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("generator", Path(__file__).with_name("generate-workbook.py"))
generator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(generator)


def build(model):
    from openpyxl import load_workbook
    path = ROOT / f"examples/{model}-example.xlsx"
    generator.build_workbook(generator.parse_template(ROOT / f"templates/{model}-template.md"), path)
    wb = load_workbook(path)
    wb.properties.creator = "jvc-analyst"
    wb.properties.lastModifiedBy = "jvc-analyst"
    wb.properties.description = "虚构数字，用于计算回归和使用示范，不是行业证据。"
    return path, wb


def add(wb, sheet, **values):
    ws = wb[sheet]
    unknown = set(values) - {c.value for c in ws[1]}
    if unknown:
        raise ValueError((sheet, unknown))
    ws.append([values.get(c.value) for c in ws[1]])


def ref(wb, sheet, field, row=2):
    headers = [c.value for c in wb[sheet][1]]
    return f"{sheet}!{get_column_letter(headers.index(field) + 1)}{row}"


def source(wb):
    fields = {c.value for c in wb["sources"][1]}
    values = dict(source_id="S1", source_name="虚构回归情景", source_type="用户假设",
                  url_or_location="examples/README.md", date="2026-10-18",
                  fields_supported="全部假设；单位：万元，市场模型除外", reliability="假设计算，不代表事实")
    add(wb, "sources", **{k: v for k, v in values.items() if k in fields})


def save(path, wb, check=None):
    from workbook_calculation import recalculate
    for ws in wb:
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == "f" or isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    wb.save(path)
    recalculate(wb, cache_output=path)
    if check:
        path.with_suffix(".checks.json").write_text(json.dumps({"schema_version": 1, "checks": [check]}, ensure_ascii=False, indent=2) + "\n")


def main():
    path, wb = build("market-sizing")
    for index, (variable, value, unit) in enumerate((("上位市场", 1000000, "元/年"), ("筛选占比", .25, "比例"), ("客户数", 100, "家"), ("点位数", 2, "个/家"), ("年价", 1000, "元/点位/年"), ("渗透率", .8, "比例")), 1):
        add(wb, "assumptions", assumption_id=f"A{index}", variable=variable, value=value, unit=unit, year=2026, geography="虚构区域", source_id="S1", confidence="假设", notes="虚构示例")
    add(wb, "top_down", line_id="T1", level="SAM", parent_line_id="A1", market_item="虚构市场", base_value="=assumptions!C2", filter_or_share="=assumptions!C3", calculated_value="=E2*F2", formula="基础市场乘筛选占比", currency="CNY", year=2026, source_id="S1")
    add(wb, "bottom_up", line_id="B1", segment="客户组甲", customer_count="=assumptions!C4", units_per_customer="=assumptions!C5", annual_price_or_spend="=assumptions!C6", penetration_rate="=assumptions!C7", calculated_value="=C2*D2*E2*F2", formula="客户乘点位乘年价乘渗透率", currency="CNY", year=2026, source_id="S1")
    add(wb, "reconciliation", metric="2026 SAM", top_down_value="=top_down!G2", bottom_up_value="=bottom_up!G2", difference="=B2-C2", difference_pct="=D2/B2", explanation="口径不确定，区间保留", preferred_value="=MIN(B2,C2)", reason="保守情景")
    add(wb, "orthogonality_check", item_a="客户组甲", item_b="无第二分组", overlap_risk="无", overlap_description="仅一个分组", resolution="唯一归属", status="clear")
    source(wb)
    save(path, wb, {"input": "assumptions!C2", "before": 1000000, "after": 1200000, "outputs": [{"cell": "top_down!G2", "before": 250000, "after": 300000}, {"cell": "reconciliation!D2", "before": 90000, "after": 140000}]})

    path, wb = build("comps-dd")
    add(wb, "companies", company_name="虚构公司甲", country_region="虚构区域", company_type="初创公司", listing_status="未上市", value_chain_position="平台", technology_route="软件", product="质检服务", target_customers="工厂", latest_year_revenue=100, revenue_year=2025, revenue_currency="CNY 万元", latest_valuation_or_market_cap=500, valuation_type="融资估值", valuation_date="2026-09-01", financing_stage="A", key_investors="未披露", comparability="收入模式可比；阶段不同", source_id="S1", confidence="假设", notes="估值单位同收入，虚构数字")
    add(wb, "segmentation", segment_type="商业模式", segment_name="按年收费", included_companies="虚构公司甲", why_it_matters="收入持续性", overlap_warning="仅一家公司", source_id="S1")
    add(wb, "coverage_notes", issue="示例不代表完整覆盖", affected_companies="全部", impact="不能用于投资判断", next_step="替换为授权材料")
    source(wb); save(path, wb)

    path, wb = build("roi-modeler")
    add(wb, "investment_terms", investment_amount=10, pre_money_valuation=90, post_money_valuation=100, initial_ownership="=A2/C2", security_type="普通股", option_pool=0, currency="CNY 万元", source_id="S1", notes="同股同权；无分红；各笔现金流均按年末发生", investment_year=2026)
    for row, year in enumerate(range(2027, 2032), 2):
        add(wb, "financial_forecast", year=year, revenue=[30,50,70,90,100][row-2], gross_margin=.5, gross_profit=f"=B{row}*C{row}", EBITDA=10, net_income=5, cash_burn=10, ending_cash=f"=K{row}-G{row}", source_id="S1", notes="现金为简化经营情景，不含未来融资；期初现金独立给定", opening_cash=100 if row==2 else f"=H{row-1}")
    add(wb, "financing_dilution", round="A", year=2027, new_money=100, pre_money_valuation=400, post_money_valuation="=C2+D2", investor_participates="否", pro_rata_amount=0, dilution_pct="=C2/E2", ownership_after_round="=(investment_terms!D2*D2+G2)/E2", source_id="S1")
    add(wb, "ownership", year_or_round=2027, ownership_start="=investment_terms!D2", dilution_pct="=financing_dilution!H2", pro_rata_investment="=financing_dilution!G2", ownership_end="=financing_dilution!I2", notes="普通股逐轮稀释")
    for row, (scenario, metric) in enumerate((("保守",50),("中性",100),("乐观",150)), 2):
        add(wb, "exit_scenarios", scenario=scenario, exit_year=2031, exit_metric="收入", exit_metric_value=metric, exit_multiple=5, enterprise_value=f"=D{row}*E{row}", net_debt_or_cash=10, equity_value=f"=F{row}-G{row}", source_id="S1", notes="净债务正数；全部为假设")
        add(wb, "returns", scenario=scenario, invested_capital="=investment_terms!A2+SUM(financing_dilution!G2:G2)", exit_year=2031, final_ownership="=ownership!E2", exit_equity_value=f"=exit_scenarios!H{row}", proceeds=f"=D{row}*E{row}", MOIC=f"=F{row}/B{row}", IRR=f"=G{row}^(1/(C{row}-investment_terms!J2))-1", key_driver="退出收入与估值倍数；此示例没有跟投")
    add(wb, "sensitivity", variable="中性退出倍数", low_case=4, base_case=5, high_case=6, impact_on_MOIC="=returns!G3", impact_on_IRR="=returns!H3")
    source(wb)
    save(path, wb, {"input": "exit_scenarios!E3", "before":5, "after":6, "outputs":[{"cell":"returns!G3","before":3.92,"after":4.72},{"cell":"returns!H3","before":3.92**.2-1,"after":4.72**.2-1}]})
    print("Built three populated synthetic workbooks and two sensitivity check files")


if __name__ == "__main__":
    main()
