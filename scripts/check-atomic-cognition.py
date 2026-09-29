#!/usr/bin/env python3
"""Exercise new profiles through the real evidence CLI using synthetic local inputs.

This checks ledger/audit behavior, not autonomous model routing or research quality.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
CORE = ROOT / "skills/jvc-research-core/scripts/researchctl.py"
INPUT = ROOT / "evals/atomic-cognition/input.md"
STAMP = "2026-09-05T00:00:00Z"
AUDIT = "jvc-claim-audit"
PLAN = "jvc-diligence-plan"


def run(*args: str, expected: int = 0) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, str(CORE), *args], cwd=ROOT.parent,
        capture_output=True, text=True, check=False,
    )
    assert result.returncode == expected, (args, result.stdout, result.stderr)
    return result


def record(skill: str, identifier: str, kind: str, **fields: object) -> dict:
    return dict(schema_version=1, record_id=identifier, record_type=kind,
                created_at=STAMP, actor="synthetic-regression", created_by_skill=skill,
                **fields)


def question(skill: str, identifier: str, **fields: object) -> dict:
    return record(skill, identifier, "question", question_text="收入记录能否解释口径差异？",
                  priority="high", hypothesis="同一期间对账能够解释差异",
                  falsifier="所谓收入是未签约管线", evidence_needed=["合同、验收与回款对账"],
                  state="gap", **fields)


def init(folder: Path, skill: str) -> None:
    folder.mkdir()
    scope = record(skill, "SC1", "scope", subject="虚构工业视觉项目",
                   decision="识别证据缺口并制定验证计划", inclusions=["提供的测试材料"],
                   exclusions=["外部核验"], geography="中国（测试场景）",
                   time_range="2026-08-20 至 2026-08-25", user_assumptions=[])
    path = folder / "scope.json"
    path.write_text(json.dumps(scope, ensure_ascii=False), encoding="utf-8")
    run("init", "--skill", skill, "--run-dir", str(folder), "--scope-file", str(path))


def append(folder: Path, records: list[dict]) -> None:
    path = folder / "input.jsonl"
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in records), encoding="utf-8")
    run("record", "--run-dir", str(folder), "--input", str(path))


def audit(folder: Path, skill: str, artifact: Path, expected: int) -> dict:
    run("audit", "--run-dir", str(folder), "--skill", skill,
        "--artifact", str(artifact), "--mode", "gate", expected=expected)
    index = json.loads((folder / "audit.json").read_text(encoding="utf-8"))
    return next(entry for entry in reversed(index["audits"]) if entry["skill"] == skill)


def check_batch2(root: Path, source_template: dict, claim_template: dict) -> None:
    fixture = ROOT / "evals/atomic-cognition/batch2-input.md"
    fixture_text = fixture.read_text(encoding="utf-8")
    # Smoke-test the shared company-claim gate, not each skill's domain behavior.
    # Domain-specific behavior is exercised separately by the batch2 text trials.
    for skill in (
        "jvc-customer-validation", "jvc-founder-assessment",
        "jvc-business-economics", "jvc-financing-milestones",
        "jvc-thesis-test", "jvc-bull-case", "jvc-bear-case",
        "jvc-industry-mechanics", "jvc-competitive-positioning",
        "jvc-return-underwriting", "jvc-decision-memo",
    ):
        folder = root / skill
        init(folder, skill)
        source = {**source_template, "created_by_skill": skill}
        claim = {**claim_template, "created_by_skill": skill,
                 "importance": "decision_critical"}
        append(folder, [source, question(skill, "Q1"), claim])
        artifact = folder / "result.md"
        artifact.write_text("研究状态：partial\n公司说法尚未核验。[S1]\n", encoding="utf-8")
        result = audit(folder, skill, artifact, 20)
        assert any(f["rule"] == "company_claim_unverified" and f["severity"] == "block"
                   for f in result["findings"]), skill
        # A label cannot make a decision-critical company assertion pass.
        assert result["status"] == "blocked", skill

    # User-specified arithmetic scenarios can be complete within their stated scope.
    for skill, text, prefix in (
        ("jvc-business-economics", "假设下单位贡献100-60-20=20万元", "所有金额单位"),
        ("jvc-financing-milestones", "给定现金情景第3月末触零，第4月末缺20万元", "现金情景以"),
    ):
        folder = root / (skill + "-scenario")
        init(folder, skill)
        source = {**source_template, "created_by_skill": skill, "location": str(fixture),
                  "source_class": "user-document", "publisher": "测试情景提供者",
                  "title": "明确给定的虚构情景", "definition": "假设计算，不是公司业绩",
                  "excerpt": next(line for line in fixture_text.splitlines() if line.startswith(prefix))}
        question_record = {**question(skill, "Q1"), "question_text": "给定假设下计算是否闭合？",
                           "hypothesis": "所有计算输入均来自给定情景", "state": "supported"}
        claim = {**claim_template, "created_by_skill": skill, "claim_text": text,
                 "claim_kind": "model_estimate", "state": "supported", "confidence": "high",
                 "importance": "decision_critical", "scope": "仅用户指定的合成情景",
                 "reasoning": "限定在提供的假设下，不宣称真实经营能力"}
        append(folder, [source, question_record, claim])
        artifact = folder / "scenario.md"
        artifact.write_text("# 假设情景\n" + text + "。[S1]\n", encoding="utf-8")
        assert audit(folder, skill, artifact, 0)["status"] == "ready"
        artifact.write_text("研究状态：partial\n" + text + "。[S1]\n", encoding="utf-8")
        wrong_status = audit(folder, skill, artifact, 20)
        assert any(f["rule"] == "artifact_status_mismatch"
                   and "expected 研究状态：ready" in f["message"] for f in wrong_status["findings"])
        artifact.write_text("# 假设情景\n" + text + "。[S999]\n", encoding="utf-8")
        missing_citation = audit(folder, skill, artifact, 20)
        assert any(f["rule"] == "artifact_source_reference" for f in missing_citation["findings"])


def check_thesis_compatibility(root: Path, source: dict, claim: dict) -> None:
    for legacy, filename in (("jvc-bull-case", "04-bull-case.md"),
                             ("jvc-bear-case", "04-bear-case.md")):
        folder = root / (legacy + "-history")
        init(folder, legacy)
        append(folder, [{**source, "created_by_skill": legacy}, question(legacy, "Q1"),
                        {**claim, "created_by_skill": legacy}])
        artifact = folder / filename
        artifact.write_text("研究状态：partial\n公司自述尚待核验。[S1]\n", encoding="utf-8")
        old_audit = audit(folder, legacy, artifact, 10)
        ledger = folder / "evidence_registry.jsonl"
        original = ledger.read_bytes()
        original_artifact = artifact.read_bytes()

        canonical = "jvc-thesis-test"
        run("init", "--skill", canonical, "--run-dir", str(folder), "--resume")
        append(folder, [question(canonical, "Q2"),
                        {**claim, "created_by_skill": canonical, "record_id": "C2",
                         "question_id": "Q2", "derived_from_claim_ids": ["C1"]}])
        unified = folder / "04-thesis-test.md"
        unified.write_text("研究状态：partial\n沿用旧材料，公司自述尚待核验。[S1]\n", encoding="utf-8")
        new_audit = audit(folder, canonical, unified, 10)
        assert new_audit["skill"] == canonical
        assert ledger.read_bytes().startswith(original), "historical records rewritten"
        index = json.loads((folder / "audit.json").read_text(encoding="utf-8"))
        assert old_audit in index["audits"], "historical audit rewritten"
        assert artifact.read_bytes() == original_artifact, "legacy artifact overwritten"
        # The original identity remains usable after the canonical entry appends.
        assert audit(folder, legacy, artifact, 10)["status"] == "partial"


def check_interview_compatibility(root: Path, source: dict, claim: dict) -> None:
    generator = ROOT / "skills/jvc-meeting-notes/scripts/generate_meeting_notes.py"
    for legacy in ("jvc-meeting-notes", "jvc-talk-notes"):
        folder = root / legacy
        init(folder, legacy)
        append(folder, [{**source, "created_by_skill": legacy}])
        artifact = folder / "【2026年08月20日访谈】虚构公司.docx"
        text = "【公司自述】收入1200万元，未经核验。[S1]"
        section = {"heading": "一、公司基本情况", "content": text}
        if legacy == "jvc-talk-notes":
            section = {"heading": "二、问答纪要", "subsections": [
                {"heading": "Q1：收入多少？", "content": "完整回答：" + text
                 + "\n对应事实层维度：财务\n待验证点：未提供凭证。"}]}
        payload = folder / "data.json"
        payload.write_text(json.dumps({"title": "虚构纪要回归", "sections": [section]},
                                      ensure_ascii=False), encoding="utf-8")
        subprocess.run([sys.executable, str(generator), str(payload), "--output", str(artifact)],
                       cwd=ROOT.parent, check=True, capture_output=True)
        # Legacy ready means source/format checks, not company-claim verification.
        old_audit = audit(folder, legacy, artifact, 0)
        ledger = folder / "evidence_registry.jsonl"
        original = ledger.read_bytes()
        original_artifact = artifact.read_bytes()
        run("init", "--skill", AUDIT, "--run-dir", str(folder), "--resume")
        append(folder, [question(AUDIT, "Q1"), claim])
        markdown = folder / "claim-audit.md"
        markdown.write_text("研究状态：partial\n" + text + "\n", encoding="utf-8")
        assert audit(folder, AUDIT, markdown, 10)["status"] == "partial"
        assert ledger.read_bytes().startswith(original)
        assert artifact.read_bytes() == original_artifact
        assert old_audit in json.loads((folder / "audit.json").read_text())["audits"]
        assert audit(folder, legacy, artifact, 0)["status"] == "ready"
        # A Word rendering does not inherit the Markdown profile's audit binding.
        rejected = audit(folder, AUDIT, artifact, 20)
        assert any(f["rule"] == "artifact_suffix" for f in rejected["findings"])
        assert audit(folder, AUDIT, markdown, 10)["status"] == "partial"


def check_structured_compatibility(root: Path, source: dict, claim: dict) -> None:
    # Synthetic records exercise audit identity and suffixes, not domain quality.
    from openpyxl import load_workbook

    for legacy, canonical, filename in (
        ("jvc-track-research", "jvc-industry-mechanics", "landscape.md"),
        ("jvc-comps-dd", "jvc-competitive-positioning", "05-comps-dd.xlsx"),
        ("jvc-roi-modeler", "jvc-return-underwriting", "05-roi-modeler.xlsx"),
        ("jvc-ic-memo", "jvc-decision-memo", "06-ic-memo-review.md"),
    ):
        folder = root / (legacy + "-migration")
        init(folder, legacy)
        for skill, suffix in ((legacy, "1"), (canonical, "2")):
            if skill == canonical:
                run("init", "--skill", skill, "--run-dir", str(folder), "--resume")
            entries = [{**source, "record_id": "S" + suffix, "created_by_skill": skill},
                       question(skill, "Q" + suffix),
                       {**claim, "record_id": "C" + suffix, "created_by_skill": skill,
                        "question_id": "Q" + suffix,
                        "support_source_ids": ["S" + suffix],
                        "derived_from_claim_ids": ["C1"] if skill == canonical else []}]
            for round_number, direction in ((1, "support"), (2, "counter")):
                entries.append(record(
                    skill, f"QU{suffix}{round_number}", "query", question_id="Q" + suffix,
                    direction=direction, query_text="合成兼容样本：收入口径的支持/反证",
                    tool_class="synthetic-fixture", target_source_class="company-material",
                    executed_at=STAMP, search_round=round_number, changed_core_judgment=False,
                    result_count=1, outcome="captured", result_summary="仅模拟检索收敛"))
            append(folder, entries)
            artifact = folder / (filename if skill == legacy else canonical[4:] + ".md")
            if artifact.suffix == ".xlsx":
                template = ROOT / "templates" / (legacy[4:] + "-template.md")
                subprocess.run([sys.executable, str(ROOT / "scripts/generate-workbook.py"),
                                str(template), str(artifact)], check=True, capture_output=True)
                workbook = load_workbook(artifact)
                workbook["sources"].append(["[S1]", "研究状态：partial；合成公司说法未核验"])
                workbook.save(artifact)
                subprocess.run([sys.executable, str(ROOT / "scripts/validate-workbook.py"),
                                str(artifact), str(template)], check=True, capture_output=True)
            else:
                artifact.write_text("研究状态：partial\n合成公司说法未核验。[S" + suffix + "]\n",
                                    encoding="utf-8")
            result = audit(folder, skill, artifact, 10)
            if skill == legacy:
                old_audit, old_artifact = result, artifact
                old_bytes = artifact.read_bytes()
                prefix = (folder / "evidence_registry.jsonl").read_bytes()
            else:
                assert (folder / "evidence_registry.jsonl").read_bytes().startswith(prefix)
                assert old_artifact.read_bytes() == old_bytes
                assert old_audit in json.loads((folder / "audit.json").read_text())["audits"]
                assert audit(folder, legacy, old_artifact, 10)["status"] == "partial"
                if old_artifact.suffix == ".xlsx":
                    rejected = audit(folder, canonical, old_artifact, 20)
                    assert any(f["rule"] == "artifact_suffix" for f in rejected["findings"])
                    assert audit(folder, canonical, artifact, 10)["status"] == "partial"


def check_return_scenario() -> dict:
    # Pure synthetic ordinary-equity scenario; not a general return calculator.
    d = Decimal
    ownership = d("0.10") * d(100) / d(125)
    proceeds = (d(300) - d(20)) * ownership
    multiple = proceeds / d(10)
    irr = multiple ** (d(1) / d(5)) - 1
    threshold = d(10) * d("1.2") ** 5 / ownership + d(20)
    debt_limit = d(300) - d(10) * d("1.2") ** 5 / ownership
    assert ownership == d("0.08") and proceeds == d("22.4") and multiple == d("2.24")
    assert abs(irr - d("0.1750317554823924")) < d("1e-16")
    assert threshold == d("331.04") and debt_limit == d("-11.04")
    assert abs(-d(10) + proceeds / (1 + irr) ** 5) < d("1e-20")
    assert (d(300) - debt_limit) * ownership == d(10) * d("1.2") ** 5
    return dict(ownership=str(ownership), proceeds=str(proceeds), moic=str(multiple),
                irr=str(irr), threshold_ev=str(threshold), net_debt_max=str(debt_limit))


def check_market_formulas(root: Path, source: dict, claim: dict) -> None:
    from openpyxl import load_workbook

    skill = "jvc-market-sizing"
    folder = root / "market-formulas"
    init(folder, skill)
    append(folder, [{**source, "created_by_skill": skill}, question(skill, "Q1"),
                    {**claim, "created_by_skill": skill, "importance": "decision_critical"}])
    artifact = folder / "market.xlsx"
    subprocess.run([sys.executable, str(ROOT / "scripts/generate-workbook.py"),
                    str(ROOT / "templates/market-sizing-template.md"), str(artifact)],
                   check=True, capture_output=True)
    workbook = load_workbook(artifact)
    workbook["sources"].append(["[S1]", "研究状态：blocked"])
    for name in ("top_down", "bottom_up", "reconciliation"):
        workbook[name]["G2"] = "文字公式不是计算公式"
    workbook.save(artifact)
    structural_negative = audit(folder, skill, artifact, 20)
    assert {f["message"].split(":")[0] for f in structural_negative["findings"]
            if f["rule"] == "artifact_workbook_formulas"} == {
                "top_down", "bottom_up", "reconciliation"}
    for name in ("top_down", "bottom_up", "reconciliation"):
        workbook[name]["G2"] = '=IF(COUNT(E2:F2)=2,E2*F2,"")'
    workbook.save(artifact)
    structural_positive = audit(folder, skill, artifact, 20)
    assert not any(f["rule"] == "artifact_workbook_formulas"
                   for f in structural_positive["findings"])

    workbook["assumptions"]["C2"] = 100
    workbook["assumptions"]["C3"] = 0.3
    workbook["top_down"]["A2"] = "TD1"
    workbook["top_down"]["E2"] = "=assumptions!C2"
    workbook["top_down"]["F2"] = "=assumptions!$C$3"
    workbook["top_down"]["G2"] = "=E2*F2"
    workbook["top_down"]["H2"] = "=E2*F2"
    workbook["bottom_up"]["A2"] = "BU1"
    for cell in ("C2", "D2", "E2", "F2"):
        workbook["bottom_up"][cell] = "=assumptions!C2"
    workbook["bottom_up"]["D2"] = 1
    workbook["bottom_up"]["G2"] = "=C2*D2*E2*F2"
    workbook["bottom_up"]["H2"] = "=C2*D2*E2*F2"
    workbook["reconciliation"]["A2"] = "TAM"
    workbook["reconciliation"]["B2"] = "=top_down!G2"
    workbook["reconciliation"]["C2"] = "=bottom_up!G2"
    workbook["reconciliation"]["D2"] = '=IF($B$2="","",$B$2-$C$2)'
    workbook["reconciliation"]["E2"] = "=IFERROR($D$2/$B$2,0)"
    workbook.save(artifact)

    # A populated row with a non-trivial literal must be blocked and identify its cell.
    workbook["top_down"]["F2"] = 0.3
    workbook.save(artifact)
    literal_top_down = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "top_down!F2" in f["message"]
               for f in literal_top_down["findings"])
    workbook["top_down"]["F2"] = "=assumptions!$C$3"

    workbook["top_down"]["F2"] = "=0.3"
    workbook["bottom_up"]["C2"] = "=100"
    workbook.save(artifact)
    literal_formula = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "top_down!F2" in f["message"]
               for f in literal_formula["findings"])
    assert any(f["rule"] == "artifact_workbook_formulas" and "bottom_up!C2" in f["message"]
               for f in literal_formula["findings"])
    workbook["top_down"]["F2"] = "=assumptions!$C$3"
    workbook["bottom_up"]["C2"] = "=assumptions!C2"

    workbook["top_down"]["G2"] = 30
    workbook.save(artifact)
    literal_calculated_value = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "top_down!G2" in f["message"]
               for f in literal_calculated_value["findings"])
    workbook["top_down"]["G2"] = "=E2*F2"

    workbook["bottom_up"]["F2"] = 0.4
    workbook.save(artifact)
    literal_bottom_up = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "bottom_up!F2" in f["message"]
               for f in literal_bottom_up["findings"])
    workbook["bottom_up"]["F2"] = "=assumptions!C2"

    workbook["reconciliation"]["D2"] = "=B3-C3"
    workbook["reconciliation"]["E2"] = "=IFERROR(D3/B3,0)"
    workbook.save(artifact)
    wrong_reconciliation = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!D2" in f["message"]
               for f in wrong_reconciliation["findings"])
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!E2" in f["message"]
               for f in wrong_reconciliation["findings"])

    workbook["reconciliation"]["D2"] = "=B2-C2+B3"
    workbook["reconciliation"]["E2"] = "=D2/B2+B3"
    workbook.save(artifact)
    extra_row_references = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!D2" in f["message"]
               for f in extra_row_references["findings"])
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!E2" in f["message"]
               for f in extra_row_references["findings"])

    workbook["reconciliation"]["D2"] = "=B2-C2+other!B2"
    workbook["reconciliation"]["E2"] = "=D2/B2"
    workbook.save(artifact)
    other_sheet_reference = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!D2" in f["message"]
               for f in other_sheet_reference["findings"])

    workbook["reconciliation"]["D2"] = "=other!$B$2-other!$C$2"
    workbook["reconciliation"]["E2"] = '="D2/B2"'
    workbook.save(artifact)
    pseudo_references = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!D2" in f["message"]
               for f in pseudo_references["findings"])
    assert any(f["rule"] == "artifact_workbook_formulas" and "reconciliation!E2" in f["message"]
               for f in pseudo_references["findings"])
    workbook["reconciliation"]["D2"] = '=IF($B$2="","",$B$2-$C$2)'
    workbook["reconciliation"]["E2"] = "=IFERROR($D$2/$B$2,0)"

    workbook["top_down"]["G2"] = "=E2*F2"
    workbook["top_down"]["H2"] = "=E2*F2+0"
    workbook.save(artifact)
    set_xlsx_formula_caches(artifact, "top_down", {"G2": "30", "H2": "30"})
    equal_cached_results = audit(folder, skill, artifact, 20)
    assert not any(f["rule"] == "artifact_workbook_formulas"
                   for f in equal_cached_results["findings"])

    workbook["top_down"]["H2"] = "=E2*F2+1"
    workbook.save(artifact)
    set_xlsx_formula_caches(artifact, "top_down", {"G2": "30", "H2": "31"})
    conflicting_results = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "top_down!H2" in f["message"]
               and "cache differs" in f["message"]
               for f in conflicting_results["findings"])

    workbook["top_down"]["H2"] = "=G2+1"
    workbook.save(artifact)
    set_xlsx_formula_caches(artifact, "top_down", {"G2": "", "H2": ""})
    missing_cache = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "top_down!H2" in f["message"]
               and "recalculate" in f["message"] for f in missing_cache["findings"])

    workbook["top_down"]["H2"] = "=bottom_up!G2"
    workbook.save(artifact)
    set_xlsx_formula_caches(artifact, "top_down", {"G2": "", "H2": ""})
    cross_sheet_cache = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "artifact_workbook_formulas" and "top_down!H2" in f["message"]
               and "recalculate" in f["message"] for f in cross_sheet_cache["findings"])

    workbook["top_down"]["H2"] = "=G2"
    workbook.save(artifact)
    positive = audit(folder, skill, artifact, 20)
    assert not any(f["rule"] == "artifact_workbook_formulas" for f in positive["findings"])
    # Formula presence does not turn missing evidence into a valid market estimate.
    assert any(f["rule"] == "company_claim_unverified" for f in positive["findings"])
    workbook.close()


def set_xlsx_formula_caches(path: Path, sheet_name: str, values: dict[str, str]) -> None:
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    doc_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    workbook = ElementTree.fromstring(parts["xl/workbook.xml"])
    relationships = ElementTree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    rel_targets = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in relationships.findall(f"{{{package_rel}}}Relationship")
    }
    sheet = next(item for item in workbook.findall(f"{{{main}}}sheets/{{{main}}}sheet")
                 if item.attrib["name"] == sheet_name)
    target = rel_targets[sheet.attrib[f"{{{doc_rel}}}id"]]
    target = target.lstrip("/") if target.startswith("/") else "xl/" + target
    sheet_root = ElementTree.fromstring(parts[target])
    for cell in sheet_root.findall(f".//{{{main}}}c"):
        ref = cell.attrib.get("r")
        if ref not in values:
            continue
        cache = cell.find(f"{{{main}}}v")
        if cache is None:
            cache = ElementTree.SubElement(cell, f"{{{main}}}v")
        cache.text = values[ref]
    parts[target] = ElementTree.tostring(sheet_root, encoding="utf-8", xml_declaration=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w") as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def check_model_input_lineage(root: Path, source: dict, claim: dict) -> None:
    skill = "jvc-business-economics"
    folder = root / "company-model-input"
    init(folder, skill)
    append(folder, [{**source, "created_by_skill": skill}, question(skill, "Q1"),
                    {**claim, "created_by_skill": skill, "claim_kind": "model_estimate",
                     "importance": "decision_critical", "state": "supported"}])
    artifact = folder / "economics.md"
    artifact.write_text("研究状态：blocked\n公司输入未经核验，仅列计算与缺口。[S1]\n",
                        encoding="utf-8")
    result = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "model_input_lineage_missing" for f in result["findings"])
    append(folder, [{**claim, "created_by_skill": skill, "record_id": "C2",
                    "importance": "material"},
                   {**claim, "created_by_skill": skill, "record_id": "C3", "supersedes": "C1",
                    "claim_kind": "model_estimate", "importance": "decision_critical",
                    "state": "supported", "derived_from_claim_ids": ["C2"]}])
    linked = audit(folder, skill, artifact, 20)
    assert not any(f["rule"] == "model_input_lineage_missing" for f in linked["findings"])
    assert any(f["rule"] == "company_claim_unverified" for f in linked["findings"])
    assert any(f["rule"] == "critical_model_input_unverified" and f["severity"] == "block"
               for f in linked["findings"])
    append(folder, [{**claim, "created_by_skill": skill, "record_id": "C4", "supersedes": "C2",
                    "claim_kind": "model_estimate", "importance": "material", "state": "supported"}])
    indirect = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "model_input_lineage_missing" and "C3 <- C4" in f["message"]
               for f in indirect["findings"])
    independent_file = folder / "independent.txt"
    independent_file.write_text("Independent synthetic evidence.", encoding="utf-8")
    append(folder, [{**source, "created_by_skill": skill, "record_id": "S2",
                    "source_class": "government", "publisher": "Independent test agency",
                    "independence_key": "synthetic-independent", "location": str(independent_file),
                    "excerpt": "Independent synthetic evidence."},
                   {**claim, "created_by_skill": skill, "record_id": "C5", "supersedes": "C4",
                    "importance": "material", "state": "supported",
                    "support_source_ids": ["S1", "S2"]}])
    verified = audit(folder, skill, artifact, 20)
    assert not any(f["rule"] in {"critical_model_input_unverified", "model_input_lineage_missing"}
                   for f in verified["findings"])
    append(folder, [{**claim, "created_by_skill": skill, "record_id": "C6", "supersedes": "C5",
                    "importance": "material", "state": "refuted",
                    "support_source_ids": ["S1", "S2"], "counter_source_ids": ["S2"],
                    "conflict_resolution": "reconciled"}])
    refuted = audit(folder, skill, artifact, 20)
    assert any(f["rule"] == "critical_model_input_unverified" and "C3 <- C6" in f["message"]
               for f in refuted["findings"])


def main() -> None:
    check_return_scenario()
    input_text = INPUT.read_text(encoding="utf-8")
    excerpt = "过去十二个月收入 1200 万元；共有 8 家客户。"
    assert excerpt in input_text
    with TemporaryDirectory(prefix="jvc-atomic-") as temporary:
        root = Path(temporary)
        folder = root / "claim"
        init(folder, AUDIT)
        source = record(AUDIT, "S1", "source", title="虚构公司融资稿", publisher="测试公司",
                        author="测试融资团队", published_at="2026-08-20", accessed_at=STAMP,
                        source_class="company-material", location=str(INPUT), excerpt=excerpt,
                        definition="收入为公司口径，未验证", geography="中国（测试场景）",
                        sample="一份虚构融资稿", statistical_scope="材料称过去十二个月",
                        stance="support", independence_key="synthetic-company")
        claim = record(AUDIT, "C1", "claim", question_id="Q1",
                       claim_text="过去十二个月收入1200万元", claim_kind="company_claim",
                       topic="revenue", importance="material", support_source_ids=["S1"],
                       counter_source_ids=[], derived_from_claim_ids=[], scope="材料自述",
                       confidence="unknown", reasoning="没有提供收入凭证",
                       conflict_resolution="none", state="unverified")
        append(folder, [source, question(AUDIT, "Q1"), claim])
        artifact = folder / "claim-audit.md"
        artifact.write_text("# 材料审查\n公司自述收入1200万元，未经核验。[S1]\n", encoding="utf-8")
        missing_label = audit(folder, AUDIT, artifact, 20)
        assert any(f["rule"] == "partial_label_missing" for f in missing_label["findings"])
        artifact.write_text("研究状态：partial\n公司自述收入1200万元，缺少凭证。[S1]\n", encoding="utf-8")
        partial = audit(folder, AUDIT, artifact, 10)
        assert partial["status"] == "partial"
        assert any(f["rule"] == "company_claim_unverified" for f in partial["findings"])

        critical = {**claim, "record_id": "C2", "supersedes": "C1", "importance": "decision_critical"}
        append(folder, [critical])
        blocked = audit(folder, AUDIT, artifact, 20)
        assert blocked["status"] == "blocked"
        assert any(f["rule"] == "company_claim_unverified" and f["severity"] == "block"
                   for f in blocked["findings"])
        assert any(f["rule"] == "artifact_status_mismatch"
                   and "expected 研究状态：blocked" in f["message"] for f in blocked["findings"])
        artifact.write_text("公司自述收入1200万元，缺少凭证。[S1]\n", encoding="utf-8")
        no_label = audit(folder, AUDIT, artifact, 20)
        assert any(f["rule"] == "blocked_label_missing" for f in no_label["findings"])
        artifact.write_text("研究状态：blocked\n公司自述收入1200万元，缺少凭证。[S1]\n",
                            encoding="utf-8")
        correct_label = audit(folder, AUDIT, artifact, 20)
        assert not any(f["rule"] in {"artifact_status_mismatch", "blocked_label_missing"}
                       for f in correct_label["findings"])
        assert any(f["rule"] == "company_claim_unverified" for f in correct_label["findings"])

        # A standalone plan can represent gaps without minting supported business claims.
        plan_folder = root / "standalone-plan"
        init(plan_folder, PLAN)
        append(plan_folder, [{**source, "created_by_skill": PLAN}, question(PLAN, "Q1")])
        plan_artifact = plan_folder / "diligence-plan.md"
        plan_artifact.write_text("研究状态：partial\n先核对收入记录；业务假设未验证。[S1]\n", encoding="utf-8")
        plan_audit = audit(plan_folder, PLAN, plan_artifact, 10)
        assert plan_audit["status"] == "partial"
        ledger = [json.loads(line) for line in (plan_folder / "evidence_registry.jsonl").read_text().splitlines()]
        assert not any(r["record_type"] == "claim" for r in ledger)

        # Deriving a business conclusion from a blocked upstream must still fail.
        run("init", "--skill", PLAN, "--run-dir", str(folder), "--resume")
        append(folder, [question(PLAN, "Q2"), record(
            PLAN, "C3", "claim", question_id="Q2", claim_text="公司已达到规模收入",
            claim_kind="agent_inference", topic="revenue", importance="material",
            support_source_ids=["S1"], counter_source_ids=[], derived_from_claim_ids=["C2"],
            scope="虚构测试", confidence="high", reasoning="错误依赖公司自述的负面测试",
            conflict_resolution="none", state="supported")])
        inherited = audit(folder, PLAN, plan_artifact, 20)
        assert any(f["rule"] == "upstream_audit_blocked" for f in inherited["findings"])

        # Neither new identity may launder the same blocked upstream as a synthesis.
        for skill, suffix in (("jvc-return-underwriting", "4"), ("jvc-decision-memo", "5")):
            run("init", "--skill", skill, "--run-dir", str(folder), "--resume")
            append(folder, [question(skill, "Q" + suffix), record(
                skill, "C" + suffix, "claim", question_id="Q" + suffix,
                claim_text="基于未核验收入，回报或投决判断已经成立",
                claim_kind="agent_inference", topic="decision", importance="decision_critical",
                support_source_ids=["S1"], counter_source_ids=[], derived_from_claim_ids=["C2"],
                scope="仅负面回归", confidence="high", reasoning="故意尝试洗掉上游阻断",
                conflict_resolution="none", state="supported")])
            rejected = audit(folder, skill, plan_artifact, 20)
            assert any(f["rule"] == "upstream_audit_blocked" for f in rejected["findings"])

        # Empty evidence cannot be disguised as a completed standalone plan.
        empty = root / "empty"
        init(empty, PLAN)
        empty_audit = audit(empty, PLAN, plan_artifact, 20)
        assert any(f["rule"] == "required_source" for f in empty_audit["findings"])
        check_batch2(root, source, claim)
        check_thesis_compatibility(root, source, claim)
        check_interview_compatibility(root, source, claim)
        check_structured_compatibility(root, source, claim)
        check_market_formulas(root, source, claim)
        check_model_input_lineage(root, source, claim)
    print("atomic cognition CLI checks passed: first-batch regressions; shared company-claim gate across eleven identities; two conditional scenarios; thesis, interview, industry, return and decision identity reuse; Word/workbook cannot inherit Markdown audit")


if __name__ == "__main__":
    main()
