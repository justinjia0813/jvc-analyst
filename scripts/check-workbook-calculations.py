#!/usr/bin/env python3
"""Regressions exercise saved XLSX files and the default delivery validator."""
import copy
import importlib.util
import json
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("validator", Path(__file__).with_name("validate-workbook.py"))
validator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(validator)


def main():
    template = ROOT / "templates/roi-modeler-template.md"
    original = ROOT / "examples/roi-modeler-example.xlsx"
    checks = ROOT / "examples/roi-modeler-example.checks.json"
    before = original.read_bytes()
    with tempfile.TemporaryDirectory(prefix="jvc-workbook-test-") as directory:
        folder = Path(directory)

        def reject(name, mutation, diagnostic):
            book = load_workbook(original)
            mutation(book)
            path = folder / (name + ".xlsx")
            book.save(path)
            saved = path.read_bytes()
            errors = validator.validate(path, template, checks_path=checks)
            assert errors and any(diagnostic in e for e in errors), (name, errors)
            assert path.read_bytes() == saved, "validation changed input"

        def blank(book):
            for ws in book:
                ws.delete_rows(2, ws.max_row)

        reject("empty", blank, "no data rows")
        reject("constant-result", lambda b: setattr(b["returns"]["G2"], "value", "=999"), "reference model inputs")
        reject("hardcoded-result", lambda b: setattr(b["returns"]["G2"], "value", 1.92), "retain a formula")
        reject("divide-zero", lambda b: setattr(b["returns"]["G2"], "value", "=F2/0"), "calculation error")
        reject("bad-reference", lambda b: setattr(b["returns"]["G2"], "value", "=#REF!/B2"), "invalid formula reference")
        reject("network-formula", lambda b: setattr(b["returns"]["G2"], "value", '=WEBSERVICE("https://example.com")+B2'), "external or invalid")
        reject("wrong-moic", lambda b: setattr(b["returns"]["G2"], "value", "=F2/B2+1"), "MOIC")
        reject("wrong-irr", lambda b: setattr(b["returns"]["H2"], "value", "=G2-1"), "IRR does not reconcile")
        reject("wrong-dilution", lambda b: setattr(b["financing_dilution"]["I2"], "value", "=investment_terms!D2"), "ownership after financing")
        reject("unknown-source", lambda b: setattr(b["investment_terms"]["H2"], "value", "S999"), "unknown source_id")
        reject("circular-reference", lambda b: setattr(b["returns"]["G2"], "value", "=G2"), "calculation error")
        assert validator.validate(original, template), "missing sensitivity file accepted"
        payload = json.loads(checks.read_text())
        payload["checks"][0]["outputs"][0]["after"] = 999
        wrong = folder / "wrong-checks.json"
        wrong.write_text(json.dumps(payload))
        errors = validator.validate(original, template, checks_path=wrong)
        assert any("sensitivity changed" in e for e in errors), errors
        # Follow-on investment at exit: use cash flows, not MOIC CAGR.
        follow = load_workbook(original)
        follow["financing_dilution"]["B2"] = 2031
        follow["financing_dilution"]["F2"] = "是"
        follow["financing_dilution"]["G2"] = 10
        follow["ownership"]["A2"] = 2031
        for row in range(2,5):
            follow["returns"][f"H{row}"] = f"=((F{row}-financing_dilution!G2)/investment_terms!A2)^(1/(C{row}-investment_terms!J2))-1"
        path = folder / "follow.xlsx"; follow.save(path)
        payload = json.loads(checks.read_text())
        payload["checks"][0]["outputs"] = [
            {"cell":"returns!G3","before":2.45,"after":2.95},
            {"cell":"returns!H3","before":3.9**.2-1,"after":4.9**.2-1}]
        valid_checks = folder / "follow.checks.json"; valid_checks.write_text(json.dumps(payload))
        assert not validator.validate(path, template, checks_path=valid_checks)
        for row in range(2,5):
            follow["returns"][f"H{row}"] = f"=G{row}^(1/(C{row}-investment_terms!J2))-1"
        follow.save(path)
        assert any("IRR does not reconcile" in e for e in validator.validate(path, template, checks_path=valid_checks))
        # A conventional intermediate follow-on must also pass with actual annual IRR.
        follow["financing_dilution"]["B2"] = 2028
        follow["ownership"]["A2"] = 2028
        for row in range(2,5):
            follow["returns"][f"H{row}"] = f"=IRR(CHOOSE({{1,2,3,4,5,6}},-investment_terms!A2,0,-financing_dilution!G2,0,0,F{row}))"
        follow.save(path)
        def annual_rate(proceeds):
            low, high = 0, 1
            for _ in range(100):
                rate = (low + high) / 2
                npv = -10 - 10 / (1 + rate)**2 + proceeds / (1 + rate)**5
                if npv > 0: low = rate
                else: high = rate
            return (low + high) / 2
        payload["checks"][0]["outputs"][1].update(before=annual_rate(49), after=annual_rate(59))
        valid_checks.write_text(json.dumps(payload))
        errors = validator.validate(path, template, checks_path=valid_checks)
        assert not errors, errors
        market = ROOT / "examples/market-sizing-example.xlsx"
        market_template = ROOT / "templates/market-sizing-template.md"
        market_checks = market.with_suffix(".checks.json")
        for name, address, formula, diagnostic in (
            ("market-wrong-arithmetic", "G2", "=E2+F2", "top_down calculation"),
            ("market-disconnected-input", "E2", "=assumptions!C2*0+1000000", "sensitivity changed"),
        ):
            book = load_workbook(market)
            book["top_down"][address] = formula
            path = folder / (name + ".xlsx"); book.save(path)
            errors = validator.validate(path, market_template, checks_path=market_checks)
            assert any(diagnostic in e for e in errors), (name, errors)
    assert original.read_bytes() == before
    print("workbook calculation regressions passed: 18 cases including follow-on IRR and sensitivity")


if __name__ == "__main__":
    main()
