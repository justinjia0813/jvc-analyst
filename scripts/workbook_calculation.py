"""Read-only delivery checks using a local LibreOffice recalculation in a temp directory."""
from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer

FUNCTIONS = {"SUM", "SUMIF", "SUMIFS", "MIN", "MAX", "ABS", "ROUND", "ROUNDUP", "ROUNDDOWN",
             "IF", "AND", "OR", "NOT", "CHOOSE", "IRR", "XIRR", "NPV", "XNPV", "POWER", "PRODUCT", "AVERAGE", "COUNT"}
FORMULA_FIELDS = {
    "market-sizing": {"top_down": ["calculated_value"], "bottom_up": ["calculated_value"],
                      "reconciliation": ["top_down_value", "bottom_up_value", "difference", "difference_pct", "preferred_value"]},
    "roi-modeler": {"investment_terms": ["initial_ownership"], "financial_forecast": ["gross_profit", "ending_cash"],
                    "financing_dilution": ["post_money_valuation", "dilution_pct", "ownership_after_round"],
                    "ownership": ["ownership_end"], "exit_scenarios": ["enterprise_value", "equity_value"],
                    "returns": ["invested_capital", "final_ownership", "exit_equity_value", "proceeds", "MOIC", "IRR"]},
    "comps-dd": {},
}


def rows(wb, name):
    ws = wb[name]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in values)]


def number(value, label):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{label}: finite calculated number required")
    return value


def equal(actual, expected, label):
    number(actual, label)
    if not math.isclose(actual, expected, rel_tol=1e-7, abs_tol=1e-7):
        raise ValueError(f"{label}: expected {expected:.10g}, got {actual:.10g}")


def static_checks(wb, model):
    for ws in wb:
        if not rows(wb, ws.title):
            raise ValueError(f"{ws.title}: no data rows; a blank template is not a deliverable")
        headers = {cell.value: cell.column for cell in ws[1]}
        for row in ws.iter_rows(min_row=2):
            if not any(c.value is not None for c in row):
                continue
            for field in FORMULA_FIELDS[model].get(ws.title, []):
                cell = row[headers[field] - 1]
                if cell.data_type != "f":
                    raise ValueError(f"{ws.title}!{cell.coordinate}: {field} must retain a formula")
                if not any(t.type == "OPERAND" and t.subtype == "RANGE" for t in Tokenizer(cell.value).items):
                    raise ValueError(f"{ws.title}!{cell.coordinate}: formula must reference model inputs")
        for row in ws:
            for cell in row:
                if cell.data_type == "e":
                    raise ValueError(f"{ws.title}!{cell.coordinate}: Excel error {cell.value}")
                if cell.data_type != "f":
                    continue
                formula = cell.value
                if any(x in formula for x in ("[", "]", "#REF!", "|", "://")):
                    raise ValueError(f"{ws.title}!{cell.coordinate}: external or invalid formula reference")
                for token in Tokenizer(formula).items:
                    if token.type == "FUNC" and token.subtype == "OPEN" and token.value[:-1].upper() not in FUNCTIONS:
                        raise ValueError(f"{ws.title}!{cell.coordinate}: unsupported function {token.value[:-1]}; review locally before extending allowlist")
    if wb._external_links:
        raise ValueError("external workbook links are not supported")
    if wb.calculation and wb.calculation.iterate:
        raise ValueError("iterative calculations are not supported")
    sources = {r["source_id"] for r in rows(wb, "sources") if r.get("source_id")}
    if not sources:
        raise ValueError("sources: source IDs required")
    for ws in wb:
        if ws.title == "sources":
            continue
        for row in rows(wb, ws.title):
            if "source_id" in row and row["source_id"] not in sources:
                raise ValueError(f"{ws.title}: missing or unknown source_id {row['source_id']!r}")


def recalculate(wb, *, cache_output=None):
    executable = os.environ.get("JVC_SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise ValueError("LibreOffice required for calculation checks; install it or set JVC_SOFFICE")
    with tempfile.TemporaryDirectory(prefix="jvc-recalc-") as folder:
        root = Path(folder)
        source, output = root / "input", root / "output"
        source.mkdir(); output.mkdir()
        wb.save(source / "model.xlsx")
        try:
            run = subprocess.run([executable, f"-env:UserInstallation={(root / 'profile').as_uri()}",
                                  "--headless", "--convert-to", "xlsx", "--outdir", str(output), str(source / "model.xlsx")],
                                 text=True, capture_output=True, timeout=45, check=False)
        except (OSError, subprocess.TimeoutExpired) as error:
            raise ValueError(f"local recalculation failed: {error}") from error
        path = output / "model.xlsx"
        if run.returncode or not path.is_file():
            raise ValueError(f"local recalculation failed: {run.stderr or run.stdout}")
        calculated = load_workbook(path, data_only=True)
        for ws in wb:
            for row in ws:
                for cell in row:
                    cached = calculated[ws.title][cell.coordinate]
                    if cached.data_type == "e" or (cell.data_type == "f" and cached.value is None):
                        raise ValueError(f"{ws.title}!{cell.coordinate}: calculation error {cached.value!r}")
        if cache_output is not None:
            Path(cache_output).write_bytes(path.read_bytes())
        return calculated


def business_checks(wb, model):
    if model == "market-sizing":
        for r in rows(wb, "top_down"):
            equal(r["calculated_value"], number(r["base_value"], "base_value") * number(r["filter_or_share"], "share"), "top_down calculation")
        for r in rows(wb, "bottom_up"):
            expected = math.prod(number(r[k], k) for k in ("customer_count", "units_per_customer", "annual_price_or_spend", "penetration_rate"))
            equal(r["calculated_value"], expected, "bottom_up calculation")
        for r in rows(wb, "reconciliation"):
            difference = r["top_down_value"] - r["bottom_up_value"]
            equal(r["difference"], difference, "reconciliation difference")
            if r["top_down_value"] == 0:
                raise ValueError("reconciliation: zero top_down_value; percentage undefined")
            equal(r["difference_pct"], difference / r["top_down_value"], "reconciliation difference_pct")
    elif model == "roi-modeler":
        terms = rows(wb, "investment_terms")
        if len(terms) != 1:
            raise ValueError("investment_terms must contain one investment")
        terms = terms[0]
        year = number(terms["investment_year"], "investment_year")
        if year != int(year):
            raise ValueError("investment_year must be an integer")
        invested = number(terms["investment_amount"], "investment_amount")
        post = number(terms["post_money_valuation"], "post_money_valuation")
        if invested <= 0 or post <= 0:
            raise ValueError("investment and post-money valuation must be positive")
        ownership = invested / post
        equal(terms["initial_ownership"], ownership, "initial_ownership")
        flows = [(year, -invested)]
        previous_year = year
        for r in rows(wb, "financing_dilution"):
            round_year = number(r["year"], "round year")
            if round_year < previous_year or round_year != int(round_year):
                raise ValueError("financing rounds must use ordered integer years")
            previous_year = round_year
            pre, new, follow = (number(r[k], k) for k in ("pre_money_valuation", "new_money", "pro_rata_amount"))
            if pre <= 0 or new < 0 or not 0 <= follow <= new:
                raise ValueError("invalid financing amounts")
            if follow > 0 and str(r["investor_participates"]).lower() not in {"是", "true", "yes", "1"}:
                raise ValueError("follow-on cash conflicts with investor_participates")
            equal(r["post_money_valuation"], pre + new, "round post-money")
            equal(r["dilution_pct"], new / (pre + new), "round dilution")
            ownership = (ownership * pre + follow) / (pre + new)
            equal(r["ownership_after_round"], ownership, "ownership after financing")
            invested += follow; flows.append((round_year, -follow))
        equal(rows(wb, "ownership")[-1]["ownership_end"], ownership, "final ownership schedule")
        scenarios = {r["scenario"]: r for r in rows(wb, "exit_scenarios")}
        returns = rows(wb, "returns")
        if len(scenarios) < 3 or len(returns) != len(scenarios) or {r["scenario"] for r in returns} != set(scenarios):
            raise ValueError("three distinct exit scenarios with matching returns required")
        for r in returns:
            label = str(r["scenario"])
            exit_year = number(r["exit_year"], "exit year")
            if exit_year <= year or exit_year < previous_year or exit_year != int(exit_year):
                raise ValueError("exit must follow investment and all financing rounds")
            equal(r["invested_capital"], invested, label + " invested capital")
            equal(r["final_ownership"], ownership, label + " final ownership")
            equal(r["exit_equity_value"], scenarios[r["scenario"]]["equity_value"], label + " equity value")
            equal(exit_year, scenarios[r["scenario"]]["exit_year"], label + " exit year")
            equal(r["MOIC"], r["proceeds"] / invested, label + " MOIC")
            if terms["security_type"] == "普通股":
                equal(r["proceeds"], r["exit_equity_value"] * ownership, label + " common-share proceeds")
            rate = number(r["IRR"], label + " IRR")
            if rate <= -1:
                raise ValueError("IRR must exceed -100%; total loss requires separate treatment")
            npv = sum(value / (1 + rate) ** (date - year) for date, value in flows)
            npv += r["proceeds"] / (1 + rate) ** (exit_year - year)
            if abs(npv) > max(1, invested) * 1e-6:
                raise ValueError(label + ": IRR does not reconcile annual investment cash flows")


def delivery_checks(path, model, checks_path=None):
    wb = load_workbook(path)
    static_checks(wb, model)
    calculated = recalculate(wb)
    business_checks(calculated, model)
    if model == "comps-dd":
        return
    if checks_path is None:
        raise ValueError("--checks required: independent baseline and changed-input expectations")
    payload = json.loads(Path(checks_path).read_text(encoding="utf-8"))
    checks = payload.get("checks")
    if payload.get("schema_version") != 1 or not isinstance(checks, list) or not checks:
        raise ValueError("checks file requires schema_version 1 and nonempty checks")

    def cell(book, ref):
        sheet, address = ref.rsplit("!", 1)
        if not re.fullmatch(r"[A-Z]+[1-9]\d*", address):
            raise ValueError("check cells must be individual A1 addresses")
        return book[sheet][address]

    for check in checks:
        mutated = load_workbook(path)
        target = cell(mutated, check["input"])
        if target.data_type == "f":
            raise ValueError("sensitivity input must be an assumption, not a formula")
        equal(target.value, number(check["before"], "input before"), "sensitivity input")
        target.value = number(check["after"], "input after")
        if check["before"] == check["after"]:
            raise ValueError("sensitivity input must change")
        outputs = check["outputs"]
        if not isinstance(outputs, list) or not outputs:
            raise ValueError("sensitivity outputs required")
        changed = recalculate(mutated)
        has_change = False
        for expected in outputs:
            if cell(wb, expected["cell"]).data_type != "f":
                raise ValueError("sensitivity output must be a formula")
            before = number(expected["before"], "expected before")
            after = number(expected["after"], "expected after")
            equal(cell(calculated, expected["cell"]).value, before, "sensitivity baseline " + expected["cell"])
            equal(cell(changed, expected["cell"]).value, after, "sensitivity changed " + expected["cell"])
            has_change |= not math.isclose(before, after, rel_tol=1e-7, abs_tol=1e-7)
        if not has_change:
            raise ValueError("sensitivity check must change at least one result")
